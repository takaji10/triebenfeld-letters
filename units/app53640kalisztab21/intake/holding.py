# -*- coding: utf-8 -*-
"""APP 53/64/0/-/Kalisz Tab.21: from the editor's fourteen scans to the edition.

    python units/app53640kalisztab21/intake/holding.py     # check the table
    ... --crop        copy the page images
    ... --xlsx        the spreadsheet for the editor's check (review/<slug>/)
    ... --take FILE   read the editor's spreadsheet back: what differs from table.csv
    ... --take FILE --apply    and write those differences into table.csv
    ... --write       page files and corpus.txt from table.csv
    ... --english     write translation/doc1.yml
    ... --cache       put it into the translation cache
    ... --summaries   write the summaries S everywhere they are kept
    ... --archive     the clean spreadsheet of the table, to keep (the unit's folder)

This file is the record of what was done to this holding. The model is AGAD
1/169/0/1/1745 (units/agad1169011745/intake/holding.py), the other holding
that is a table read from the scans.

The volume is the register of the hearth tax (podymne) collected in the
Kalisz voivodeship for the St John's instalment of 1789, on a printed form.
The editor has the cover and thirteen pages of it (2026-10-13): leaf 1,
leaves 7 to 9 and leaves 18 to 20, recto and verso. There was no
transcription. The table was read from the scans in the working session
and is kept in table.csv beside this file, one line for each line of the
register: that file is the reading, and everything else is made from it.

The file named "017v.jpg" is the verso of leaf 9, not of leaf 17. Leaf 9
recto ends with entry 75 of the Kalisz district and this page runs from 76
to 107, with places of the Kalisz district; leaf 18 recto begins at entry 58
of another district, so the page before it would end at 57. A verso carries
no leaf number, so the file's name is the only thing that said 17. Its page
id is 0009_b.

Page ids are the leaf number and the side: 0001_a; 0007_a, 0007_b; 0008_a,
0008_b; 0009_a, 0009_b; 0018_a, 0018_b; 0019_a, 0019_b; 0020_a, 0020_b. The
cover is 0000_a and stands before the document as front matter.

What is on the pages. Each page is a printed form: "Kwity. Nro." (the
number of the receipt), "Nazwiska Miast i Wsiow." (the names of the towns
and villages), "Summa wybrana." (the sum collected) and "Summa Delaty."
(the sum of the delata, an order to pay the tax straight to the army), each
sum in "Złote." and "Gr:". Leaf 1 has the title above the form. A place
has a figure under one of the two sums. The receipts are numbered district
by district: "Powiat Pyzdrski" (leaf 1, entries 1 to 29; leaves 7 and 8,
entries 379 to 448), "Powiat Kaliski" (leaves 8 and 9, entries 1 to 107),
and a third district whose heading is on none of these pages (leaves 18 to
20, entries 58 to 238): its places are those of the Konin district.

How the table is written down (notes.md has the reasons):
- the heading of the form is given once, at the head of the first page;
- a cell the clerk left empty is empty; the stop after a figure is left out;
- a name is given as written, with the marks that are on the page and no
  others; a word the clerk shortened keeps his colon ("Pleban:", "Częś:");
- a name not read with certainty carries "[?]", or stands as "[Tra…ry?]"
  where only some letters are read, and table.csv says in its last column
  what the doubt is.

A check built in. The pages carry no totals, so nothing adds up. But the tax
was laid on hearths at a fixed rate, and the sums show it: in the Pyzdry
and Kalisz districts every village's sum is a multiple of 3 złoty 15 groszy
(1 złoty is 30 groszy), and on leaves 18 to 20 a multiple of 2 złoty 15
groszy. check() lists every line that is not; a town may have another
rate. A misread digit usually breaks the multiple.

The pictures in the spreadsheet. The lines of a page are ruled at even
distances, so the place of each line is worked out from the first and the
last line of its page (POS), which were read off the scans. The place is
only near enough: a picture shows two or three lines, and the editor finds
the line by its number. Nothing in the edition depends on it.
"""
import csv
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
UNIT_DIR = os.path.dirname(HERE)
ROOT = os.path.dirname(os.path.dirname(UNIT_DIR))
sys.path.insert(0, os.path.join(ROOT, 'pipeline', 'intake'))
import courtbook  # noqa: E402

SLUG = 'app53640kalisztab21'
REF = 'APP 53/64/0/-/Kalisz Tab.21'
RAW = (r"J:\Documents\Archive\Genealogy\References\Poznań State Archives"
       r"\Podymne z woj. kaliskiego [protocollon] 1789, 53.64.0.-.Kalisz Tab.21)")

# page label in table.csv -> (file, page id)
PAGES = [('1', '001.jpg', '0001_a'), ('7', '007.jpg', '0007_a'), ('7v', '007v.jpg', '0007_b'),
         ('8', '008.jpg', '0008_a'), ('8v', '008v.jpg', '0008_b'), ('9', '009.jpg', '0009_a'),
         ('9v', '017v.jpg', '0009_b'), ('18', '018.jpg', '0018_a'), ('18v', '018v.jpg', '0018_b'),
         ('19', '019.jpg', '0019_a'), ('19v', '019v.jpg', '0019_b'), ('20', '020.jpg', '0020_a'),
         ('20v', '020v.jpg', '0020_b')]
PAGE_ID = {p: pid for p, _f, pid in PAGES}
FILE = {p: f for p, f, _pid in PAGES}
SCANS = [('000.jpg', None, '0000_a', None)] + [(f, None, pid, None) for _p, f, pid in PAGES]
DOCS = [(1, [pid for _p, _f, pid in PAGES])]
# leaves 18 to 20: the third district, hearths at 2 zł 15 gr
THIRD = ('18', '18v', '19', '19v', '20', '20v')

COVER = ('0000_a', [
    'Podymne',
    'z Wojewodz: Kalisk:',
    'Raty Sto Jańsk:',
    '1789.',
    'Lib: 24. Hub[?]: 22.',
    'TAB. KALISZ 21',
])
TITLE = ['REGESTR',
         'Wybranego Podatku Podymnego i z tegoż Podatku Wydanych Delat Woysku Koronnemu pod '
         'Tytułem Offiara na ten raz ustanowiona w Racie Swięto-Jańskiey Roku 1789.']
TITLE_EN = ['REGISTER',
            'of the hearth tax collected, and of the orders to pay issued out of that tax to the Crown army, under '
            'the title of the Offering established for this occasion, in the St John\'s instalment of the year 1789.']
# The printed title stands over the name column of the first heading row:
# the site shows it as the caption of the table (unitlib.table_of). The
# labels of the first two columns stand in the last heading row.
HEAD = [['', ' '.join(TITLE), 'Summa wybrana.', '', 'Summa Delaty.', ''],
        ['Kwity. Nro.', 'Nazwiska Miast i Wsiow.', 'Złote.', 'Gr:', 'Złote.', 'Gr:']]
HEAD_EN = [['', ' '.join(TITLE_EN), 'Sum collected.', '', 'Sum of the order to pay.', ''],
           ['Receipts. No.', 'Names of the towns and villages.', 'Złoty.', 'Gr.', 'Złoty.', 'Gr.']]
FIGURES = ['w_zl', 'w_gr', 'd_zl', 'd_gr']

# The English of the common words; names stand as written, except the
# edition's settled forms (NAME_EN). Applied in this order.
WORDS = [
    (r'^Powiat Pyzdrski$', 'Pyzdry district'), (r'^Powiat Kaliski$', 'Kalisz district'),
    (r'\bMiasteczko\b', 'Small town'), (r'\bMiasto\b', 'Town'), (r'^Wie[sś]\b(?! Kuny)', 'Village'),
    (r'^Wies Kuny z Młynem$', 'Village Kuny with the mill'),
    (r'\bTamże [Pp]oddanych Pleb(?:an)?:', 'The same place, subjects of the parish priest'),
    (r'\bTamże Poddani Pleban?a?:', 'The same place, subjects of the parish priest'),
    (r'\bTamże Poddanych\b', 'The same place, the subjects'),
    (r'\bTamże Plebanskich\b', 'The same place, of the parish priest'),
    (r'^Plebanskich$', 'Of the parish priest'),
    (r'\bTamże Woytostwo\b', 'The same place, the wójt\'s holding'),
    (r'\bTamże Wybraństwo\b', 'The same place, the wybraniec\'s holding'),
    (r'^Woytostwo Pyzdrskie$', 'Wójt\'s holding of Pyzdry'), (r'^Zamek Pyzdrski$', 'Pyzdry castle'),
    (r'\bWoytostwo\b', 'Wójt\'s holding'), (r'\bPlebania w\b', 'Parsonage at'),
    (r'\bz Folwarkiem\b', 'with the demesne farm'), (r'\bFolwark\b', 'demesne farm'),
    (r'\bTrzy Folwarki\b', 'Three demesne farms'),
    (r'\bCzęść do Kościoła\b', 'part belonging to the church of'),
    (r'\bz Części w\b', 'from the part at'), (r'\bz Częś(?:ci|:)', 'from the part of'),
    (r'\bCzęść Wsi\b', 'Part of the village of'), (r'\bCzęści\b', 'parts of'), (r'\bCzęść do\b', 'part belonging to'),
    (r'\bCzęść\b', 'part of'),
    (r'\bPustkowie\b', 'Isolated farm'), (r'\bz Pustk\[uncertain: owie\]: Młynem', 'with the [uncertain: isolated farm] and the mill'),
    (r'\bz Młyne(?:m|:)', 'with the mill'), (r'\bMłyn\b', 'Mill'),
    (r'\bO[lł][eę]n?dry\b', 'Hollander settlement'), (r'\bHolendry\b', 'Hollander settlement'),
    (r'\bGrunta Dziedziczne\b', 'Hereditary lands'), (r'\bDzierzawa\b', 'Lease'), (r'\bZamek\b', 'Castle'),
    (r'\bw Par:', 'in the parish of'), (r' y ', ' and '), (r'\bJP\.', 'Mr'),
    (r' z (?=[A-ZĆŁŚŻ])', ' with '),
]
# "Nowa Wies" (182, between Osiny and the Trąbczyn Olęders) is the edition's
# Nowawies; the "Nowa Wieś" of entries 430 and 118 are other villages.
NAME_EN = (('Trąpczyn', 'Trąbczyn'), ('Trąbczynskie', 'of Trąbczyn'), ('Nowa Wies', 'Nowawies'))


def table():
    with io.open(os.path.join(HERE, 'table.csv'), encoding='utf-8', newline='') as f:
        rows = list(csv.DictReader(f))
    for r in rows:                      # the label in the reading is the file's: "17v" is leaf 9 verso
        if r['page'] == '17v':
            r['page'] = '9v'
    return rows


def money(r, a):
    """A sum in groszy, or None where the cells are empty."""
    zl, gr = r[a + '_zl'], r[a + '_gr']
    if not zl and not gr:
        return None
    return int(zl or 0) * 30 + int(gr or 0)


def check(rows):
    """What does not hold: (problems, lines to look at)."""
    bad, look = [], []
    run = None
    for r in rows:
        if not r['no']:
            run = None
            continue
        n = int(r['no'])
        if run is not None and n != run + 1 and r['page'] not in ('7', '18'):
            bad.append('no. %s on leaf %s does not follow %d' % (r['no'], r['page'], run))
        run = n
        w, d = money(r, 'w'), money(r, 'd')
        if (w is None) == (d is None):
            bad.append('leaf %s, no. %s: a figure is expected under one of the two sums' % (r['page'], r['no']))
            continue
        unit = 75 if r['page'] in THIRD else 105
        if (w if d is None else d) % unit:
            look.append((r['page'], r['no'], r['name'], w if d is None else d))
    return bad, look


def line(cells):
    out = io.StringIO()
    csv.writer(out, lineterminator='').writerow(cells)
    return out.getvalue()


def read_source():
    """{(1, page id): [the lines of the page, as CSV]}: the title and the heading open the first page."""
    parts = {}
    for r in table():
        key = (1, PAGE_ID[r['page']])
        if not parts:
            parts[key] = [line(h) for h in HEAD]
        parts.setdefault(key, []).append(line([r['no'], r['name']] + [r[k] for k in FIGURES]))
    return parts


def en_words(s):
    s = re.sub(r'\[([^\]\[]+)\?\]', r'[uncertain: \1]', re.sub(r'(\S+?)\[\?\]', r'[uncertain: \1]', s))
    for pat, new in WORDS:
        s = re.sub(pat, new, s)
    for old, new in NAME_EN:
        s = s.replace(old, new)
    return s[:1].upper() + s[1:]


def english():
    pages = {}
    for r in table():
        pid = PAGE_ID[r['page']]
        if not pages:
            pages[pid] = [line(h) for h in HEAD_EN]
        pages.setdefault(pid, []).append(line([r['no'], en_words(r['name'])] + [r[k] for k in FIGURES]))
    return {1: ['\n'.join(pages[pid]) for pid in DOCS[0][1]]}


# ---------------------------------------------------------------- the spreadsheet for the editor's check

# page -> (where its first and its last line stand in the reading strips: strip number, y at a width of 2000)
POS = {'1': ((1, 361), (4, 408)), '7': ((1, 85), (4, 455)), '7v': ((1, 95), (4, 465)), '8': ((1, 120), (4, 470)),
       '8v': ((1, 140), (4, 545)), '9': ((1, 120), (4, 505)), '9v': ((1, 135), (4, 495)), '18': ((1, 90), (4, 520)),
       '18v': ((1, 65), (4, 505)), '19': ((1, 75), (4, 520)), '19v': ((1, 80), (4, 520)), '20': ((1, 95), (4, 570)),
       '20v': ((1, 95), (4, 520))}
XLSX = 'table_check.xlsx'
FIRST = 6


def row_pictures(out_dir, rows, scale=0.5, half=125):
    from PIL import Image
    os.makedirs(out_dir, exist_ok=True)
    made = {}
    for page, _f, _pid in PAGES:
        mine = [r for r in rows if r['page'] == page]
        im = Image.open(os.path.join(RAW, FILE[page])).convert('RGB')
        W, H = im.size
        y0, y1 = int(H * 0.115), int(H * 0.965)
        step = (y1 - y0) // 4
        wide = (int(W * 0.56) - int(W * 0.03)) + (int(W * 0.94) - int(W * 0.585)) + 6

        def full(pos):
            k, yd = pos
            return y0 + max(0, (k - 1) * step - 60) + yd * wide / 2000.0
        a, b = full(POS[page][0]), full(POS[page][1])
        for i, r in enumerate(mine):
            y = a + (b - a) * i / max(1, len(mine) - 1)
            c = im.crop((int(W * 0.02), int(y - half - 15), int(W * 0.96), int(y + half - 15)))
            c = c.resize((int(c.width * scale), int(c.height * scale)), Image.LANCZOS)
            path = os.path.join(out_dir, '%s_%03d.jpg' % (page, i))
            c.save(path, quality=80)
            made[(page, i)] = path
    return made


def xlsx(rows):
    path = os.path.join(ROOT, 'review', SLUG, XLSX)
    if os.path.exists(path) and '--again' not in sys.argv:
        sys.exit('%s is there, and may be the file the editor corrected. Move it away, or add --again.' % path)
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    rdir = os.path.join(ROOT, 'review', SLUG)
    pics = row_pictures(os.path.join(rdir, 'row_img'), rows)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Register'
    bold = Font(bold=True)
    ws['A1'] = '%s, the hearth-tax register of 1789: leaves 1, 7 to 9 and 18 to 20' % REF
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = ('Each line of the register is one line here, with a picture of that part of the scan at the right: the '
                'picture shows two or three lines, so find the line by its number at the left of the picture. Correct a cell '
                'where the scan has something else, clear a doubt you have settled, and write anything more under '
                '"Your remark". Save the file and tell Claude.')
    ws['A3'] = ('"[?]" marks a name not read with certainty; "Doubt" says why. The pages have no totals, so the figures '
                'cannot be added up. "Rate?" says LOOK where a sum is not a multiple of the rate of its district '
                '(3 zł 15 gr; on leaves 18 to 20, 2 zł 15 gr): look at those figures first. Towns may have another rate.')
    heads = ['Leaf', 'No.', 'Name (Nazwiska Miast i Wsiow)', 'Collected: złote', 'gr', 'Delata: złote', 'gr',
             'Rate?', 'Doubt (Claude)', 'Your remark', 'The line on the scan']
    for j, t in enumerate(heads, 1):
        c = ws.cell(FIRST - 1, j, t)
        c.font, c.alignment = bold, Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(bottom=Side(style='medium'))
    ws.row_dimensions[FIRST - 1].height = 34
    doubt = PatternFill('solid', fgColor='FFF2CC')
    _bad, look = check(rows)
    looks = {(p, n) for p, n, _name, _v in look}
    count = {}
    n = FIRST
    for r in rows:
        i = count.get(r['page'], 0)
        count[r['page']] = i + 1
        ws.cell(n, 1, r['page'])
        if r['no']:
            ws.cell(n, 2, int(r['no']))
        ws.cell(n, 3, r['name'])
        for j, k in enumerate(FIGURES):
            if r[k]:
                ws.cell(n, 4 + j, int(r[k])).alignment = Alignment(horizontal='right', vertical='center')
        if not r['no']:
            ws.cell(n, 3).font = bold
        if (r['page'], r['no']) in looks:
            c = ws.cell(n, 8, 'LOOK')
            c.font, c.fill = bold, PatternFill('solid', fgColor='FF7C80')
        if r['note']:
            c = ws.cell(n, 9, r['note'])
            c.fill = doubt
            if '[' in r['name']:
                ws.cell(n, 3).fill = doubt
        pic = XLImage(pics[(r['page'], i)])
        ws.add_image(pic, 'K%d' % n)
        ws.row_dimensions[n].height = pic.height * 0.75 + 2
        for j in range(1, 11):
            ws.cell(n, j).alignment = Alignment(horizontal=ws.cell(n, j).alignment.horizontal, vertical='center',
                                                wrap_text=(j in (3, 9, 10)))
        n += 1
    for j, wd in enumerate([6, 6, 38, 10, 6, 10, 6, 8, 44, 36, 200], 1):
        ws.column_dimensions[L(j)].width = wd
    ws.freeze_panes = 'D%d' % FIRST
    wb.save(path)
    print('wrote', path, 'with', n - FIRST, 'lines')


def take(path, apply):
    """The editor's spreadsheet against table.csv: what they changed, and what they wrote."""
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    rows = table()
    got = []

    def s(x):
        return '' if x is None else str(x).strip()
    for n in range(FIRST, ws.max_row + 1):
        v = [ws.cell(n, j).value for j in range(1, 11)]
        if v[0] is None and v[2] is None:
            continue
        fig = ['' if s(x) == '' else str(int(float(s(x)))) for x in v[3:7]]
        got.append(dict(zip(['page', 'no', 'name'] + FIGURES + ['note', 'remark'],
                            [s(v[0]), s(v[1]), s(v[2])] + fig + [s(v[8]), s(v[9])])))
    assert len(got) == len(rows), 'the spreadsheet has %d lines, table.csv %d' % (len(got), len(rows))
    changed = notes = 0
    for old, new in zip(rows, got):
        label = '%s/%s' % (old['page'], old['no'] or old['name'])
        for k in ['no', 'name'] + FIGURES:
            if old[k] != new[k]:
                print('  %-12s %-6s %r -> %r' % (label, k, old[k], new[k]))
                changed += 1
                if apply:
                    old[k] = new[k]
        if new['remark']:
            print('  %-12s remark %s' % (label, new['remark']))
        if old['note'] != new['note']:
            notes += 1
            if apply:
                old['note'] = new['note']
    print(changed, 'cell(s) differ from table.csv;', notes, 'note(s) cleared or changed in the "Doubt" column')
    if apply and (changed or notes):
        with io.open(os.path.join(HERE, 'table.csv'), 'w', encoding='utf-8', newline='') as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator='\n')
            w.writeheader()
            w.writerows(rows)
        print('table.csv written. Look at every changed name that lost its last character against the scan; then '
              '--write, --english, --cache, --archive and rebuild.')


ARCHIVE = 'APP_53_64_0_Kalisz_Tab_21_register.xlsx'


def archive(rows):
    """The table to keep: as written, in English, and a sheet on the source. Values only. After the editor's check."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter as L
    bold, grey = Font(bold=True), Font(italic=True, color='666666')
    wb = Workbook()

    def sheet(ws, head, words, title):
        ws['A1'], ws['A2'] = title[0], title[1]
        ws['A1'].font = Font(bold=True, size=14)
        ws['A3'] = 'Archiwum Państwowe w Poznaniu, 53/64/0/-/Kalisz Tab.21 (Zbiór tabel podatkowych i lustracji)'
        for i, h in enumerate(head):
            for j, t in enumerate(h):
                if i == 0 and j == 1:           # the title, given above
                    continue
                c = ws.cell(5 + i, 2 + j, t)
                c.font, c.alignment = bold, Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(6, 1, 'Leaf').font = grey
        ws.cell(6, 8, 'Note').font = grey
        ws.merge_cells('D5:E5')
        ws.merge_cells('F5:G5')
        for j in range(1, 9):
            ws.cell(6, j).border = Border(bottom=Side(style='medium'))
        n = 7
        for r in rows:
            ws.cell(n, 1, r['page']).font = grey
            if r['no']:
                ws.cell(n, 2, int(r['no']))
            ws.cell(n, 3, words(r['name']))
            if not r['no']:
                ws.cell(n, 3).font = bold
            for j, k in enumerate(FIGURES):
                if r[k]:
                    ws.cell(n, 4 + j, int(r[k]))
            if r['note']:
                ws.cell(n, 8, r['note']).font = grey
            n += 1
        for j, wd in enumerate([6, 8, 44, 10, 6, 10, 6, 70], 1):
            ws.column_dimensions[L(j)].width = wd
        ws.freeze_panes = 'D7'

    ws = wb.active
    ws.title = 'Register'
    sheet(ws, HEAD, lambda s: s, TITLE)
    sheet(wb.create_sheet('English'), HEAD_EN, en_words, TITLE_EN)
    path = os.path.join(UNIT_DIR, ARCHIVE)
    wb.save(path)
    print('wrote', path)


# ---------------------------------------------------------------- the summaries

S = {
 1: ('Dreizehn Seiten des Registers der Rauchfangsteuer (podymne) der Woiwodschaft Kalisch für die Johannisrate 1789, auf einem gedruckten Formular in polnischer Sprache. Der gedruckte Titel nennt es das Register der erhobenen Rauchfangsteuer und der daraus dem Kronheer erteilten Zahlungsanweisungen (delaty). Jede Zeile gibt die Nummer der Quittung, den Namen einer Stadt, eines Dorfes oder eines Dorfanteils und eine Summe in Złoty und Groschen, entweder als erhoben oder als Anweisung. Die Seiten enthalten den Kreis Pyzdry mit den Einträgen 1 bis 29 und 379 bis 448, den Kreis Kalisch mit den Einträgen 1 bis 107 und die Einträge 58 bis 238 eines dritten Kreises, dessen Überschrift auf diesen Seiten fehlt; seine Orte, darunter Tuliszkow, Władysławow und Rychwał, liegen um Konin. Trąbczyn, geschrieben „Trąpczyn“, ist dort Eintrag 178 mit 80 Złoty. Es folgen die Untertanen des Pfarrers am selben Ort mit 20 Złoty, Łazy mit 12 Złoty 15 Groschen, Osiny mit 17 Złoty 15 Groschen, Nowa Wieś mit 12 Złoty 15 Groschen und die Trąbczyner Olęder („Olendry Trąbczynskie“) mit 35 Złoty. Łukom steht als Eintrag 106 mit 97 Złoty 15 Groschen unter den Anweisungen, die Olęder-Siedlung „Olędry Lusnie“ als Eintrag 109 mit 35 Złoty und Szetlewek als Eintrag 172 mit 55 Złoty. Summen der Seiten oder der Kreise stehen auf diesen Seiten nicht.',
     'Thirteen pages of the register of the hearth tax (podymne) of the Kalisz voivodeship for the St John\'s instalment of 1789, on a printed form, in Polish. The printed title calls it the register of the hearth tax collected and of the orders to pay (delaty) issued out of it to the Crown army. Each line gives the number of the receipt, the name of a town, a village or a part of a village, and a sum in złoty and groszy, either as collected or as an order to pay. The pages hold the Pyzdry district with entries 1 to 29 and 379 to 448, the Kalisz district with entries 1 to 107, and entries 58 to 238 of a third district whose heading is not on these pages; its places, among them Tuliszkow, Władysławow and Rychwał, lie around Konin. Trąbczyn, written "Trąpczyn", is entry 178 there, with 80 złoty. It is followed by the subjects of the parish priest at the same place with 20 złoty, Łazy with 12 złoty 15 groszy, Osiny with 17 złoty 15 groszy, Nowa Wieś with 12 złoty 15 groszy, and the Trąbczyn Olęders ("Olendry Trąbczynskie") with 35 złoty. Łukom is entry 106, with 97 złoty 15 groszy among the orders to pay; the Olęder settlement "Olędry Lusnie" is entry 109 with 35 złoty, and Szetlewek entry 172 with 55 złoty. These pages have no totals for a page or a district.'),
}
HOW =('written in the working session from the table as read (intake/holding.py, 2026-10-13).')


def with_cover(fn):
    """Write the document, then put the cover before it as front matter (as in APP 53/968/0/-/801)."""
    fn()
    path = os.path.join(UNIT_DIR, 'corpus.txt')
    head = '[PAGE %s]\n%s\n' % (COVER[0], '\n'.join(COVER[1]))
    s = io.open(path, encoding='utf-8').read()
    if not s.startswith(head):
        io.open(path, 'w', encoding='utf-8', newline='\n').write(head + s)
    io.open(os.path.join(UNIT_DIR, 'transcriptions', COVER[0] + '.txt'), 'w', encoding='utf-8',
            newline='\n').write('\n'.join(COVER[1]) + '\n')


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    arg = sys.argv[1:]
    rows = table()
    bad, look = check(rows)
    print(len(rows), 'lines;', sum(1 for r in rows if r['no']), 'entries;', sum(1 for r in rows if '?]' in r['name']),
          'names in doubt;', sum(1 for r in rows if r['note']), 'with a note;', len(look), 'sums off the rate;',
          'PROBLEMS' if bad else 'numbers run on')
    for b in bad:
        print('  ', b)
    if '--look' in arg:
        for p, n, name, v in look:
            print('   leaf %-4s no. %-4s %-34s %d zł %d gr' % (p, n, name, v // 30, v % 30))
    if '--crop' in arg:
        courtbook.cut(RAW, SCANS)
    if '--xlsx' in arg:
        xlsx(rows)
    if '--take' in arg:
        take(arg[arg.index('--take') + 1], '--apply' in arg)
    if '--archive' in arg:
        archive(rows)
    if '--write' in arg:
        assert not bad, 'the table has problems'
        with_cover(lambda: courtbook.write_pages(UNIT_DIR, DOCS, read_source()))
    if '--english' in arg:
        tdir = os.path.join(HERE, 'translation')
        os.makedirs(tdir, exist_ok=True)
        courtbook.write_doc_yml(os.path.join(tdir, 'doc1.yml'), english()[1])
    if '--cache' in arg:
        courtbook.save_cache(ROOT, SLUG, os.path.join(HERE, 'translation'))
    if '--summaries' in arg:
        courtbook.write_summaries(ROOT, UNIT_DIR, SLUG, REF, S, HOW)


if __name__ == '__main__':
    main()
