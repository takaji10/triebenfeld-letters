# -*- coding: utf-8 -*-
"""AGAD 1/169/0/1/1745: from the editor's three scans to the edition.

    python units/agad1169011745/intake/holding.py            # check the table
    ... --crop        cut the page images
    ... --sheet       the fold page for the editor
    ... --xlsx        the spreadsheet for the editor's check (review/<slug>/)
    ... --take FILE   read the editor's spreadsheet back: what differs from table.csv
    ... --take FILE --apply    and write those differences into table.csv
    ... --write       page files and corpus.txt from table.csv
    ... --english     write translation/doc1.yml
    ... --cache       put it into the translation cache
    ... --summaries   write the summaries S everywhere they are kept

This file is the record of what was done to this holding.

The volume is a table of the villages of a district of South Prussia. The
editor has three scans, each an opening: six pages, numbered 401 to 406 in
pencil, with the entries 140 to 303 (2026-10-12). The third file is named
"404-405.jpg" and shows pages 405 and 406. Page ids carry the pencil number
and the side of the opening: 0401_a1, 0402_a2, 0403_a1, 0404_a2, 0405_a1,
0406_a2.

There was no transcription. The table was read from the scans in the
working session and is kept in table.csv beside this file, one line for
each line of the table: that file is the reading, and corpus.txt, the page
files, the spreadsheet and the English are all made from it. A change to a
name or a figure is made in table.csv (by hand, or with --take from the
spreadsheet the editor has corrected), and then --write, --english and
--cache are run again and the holding rebuilt.

The heading is printed, on a taller leaf that lies behind the written
leaves and shows above each of them (pencil 397 on the left, a struck-out
408 on the right): one heading serves every opening. It is given once, at
the head of the first page, as HEAD.

How the table is written down (notes.md has the reasons):
- the kind of place (Dorf, Haul., Mühle, Vorw., Kloster, Krug), in German
  script, and the name, in Latin script, stand in one column of the page and
  in one cell of the text; table.csv and the spreadsheet keep them apart so
  that each can be checked;
- the stop after every figure and number is left out; the clerk's dash in
  an empty column is "—";
- the tailed stroke with which the clerk ends a shortened name is given as
  a period, or as the colon where he set one after it;
- a name not read with certainty carries "[?]", and table.csv says in its
  last column what the doubt is.

The check built in: in every line the men and women add up to the sum and
the sums to the total, and the lines of each page add up to its "Latus" in
all nine columns. All of it holds as read, which is the witness for the
figures; the names have no such witness.

rows.json has, for every line, where it stands on its scan (the height of
the line and the slant of the page). It is only used to cut the picture of
each line that the spreadsheet shows; nothing in the edition depends on it.

The folds are read off strips of the gutter; first proposals for the editor.
"""
import csv
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
UNIT_DIR = os.path.dirname(HERE)
ROOT = os.path.dirname(os.path.dirname(UNIT_DIR))
sys.path.insert(0, os.path.join(ROOT, 'pipeline', 'intake'))
import courtbook  # noqa: E402

SLUG = 'agad1169011745'
REF = 'AGAD 1/169/0/1/1745'
RAW = (r"J:\Documents\Archive\Genealogy\References\Central Archives of Historical Records in Warsaw (AGAD)"
       r"\Die neuesten topografischen Nachrichten von Südpreussen (1.169.0.1.1745) (1797)")

# (file, x of the fold, left page, right page). First proposals.
SCANS = [
    ('401-402.jpg', 3122, '0401_a1', '0402_a2'),
    ('403-404.jpg', 2625, '0403_a1', '0404_a2'),
    ('404-405.jpg', 2632, '0405_a1', '0406_a2'),
]
PAGE_ID = {'401': '0401_a1', '402': '0402_a2', '403': '0403_a1', '404': '0404_a2', '405': '0405_a1', '406': '0406_a2'}
LEAF = {v: 'page ' + k for k, v in PAGE_ID.items()}
DOCS = [(1, [PAGE_ID[p] for p in sorted(PAGE_ID)])]
FOLD_NOTE = 'Both halves are pages of the edition.'

FIGURES = ['houses', 'churches', 'chr_m', 'chr_f', 'chr_sum', 'jew_m', 'jew_f', 'jew_sum', 'total']

# The printed heading, level by level. A label stands over the blanks that
# follow it; "männlichen" and "weiblichen" share the one word "Geschlechts"
# on the page, which is given to each.
HEAD = [
    ['', '', 'Anzahl der Häuser', 'Anzahl der Kirchen', 'Menschen-Zahl', '', '', '', '', '', 'Summa aller Menschen'],
    ['', '', '', '', 'Christen', '', '', 'Juden', '', '', ''],
    ['No.', 'Nahmen der Dörfer', '', '', 'männlichen Geschlechts', 'weiblichen Geschlechts', 'Summa Christen',
     'männlichen Geschlechts', 'weiblichen Geschlechts', 'Summa', ''],
]
HEAD_EN = [
    ['', '', 'Number of houses', 'Number of churches', 'Population', '', '', '', '', '', 'Total of all persons'],
    ['', '', '', '', 'Christians', '', '', 'Jews', '', '', ''],
    ['No.', 'Names of the villages', '', '', 'male', 'female', 'Total Christians', 'male', 'female', 'Total', ''],
]
# The kinds of place, as the English gives them. "Hollander" is the
# edition's word for a Hauländer (reference/translation_glossary.yml).
KIND_EN = {'Dorf': 'Village', 'Haul.': 'Hollander settlement', 'Alt Haul.': 'Old Hollander settlement',
           'Mühle': 'Mill', 'Vorw.': 'Demesne farm', 'Kloster': 'Monastery', 'Krug': 'Inn', 'd°': 'ditto', '': ''}
# Names the edition has settled take its form in the English; every other
# name stands as the clerk wrote it. "Nowa wieś Bru." (192) is not the
# edition's Nowawies, which is entry 194, and stays as written.
NAME_EN = (('Trąpczyn', 'Trąbczyn'), ('Swiątniki', 'Świątniki'))


def table():
    with io.open(os.path.join(HERE, 'table.csv'), encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))


def num(v):
    return 0 if v in ('-', '') else int(v)


def check(rows):
    """Every line adds up, every page adds up to its Latus, the numbers run on. Returns the problems."""
    bad = []
    nos = [int(r['no']) for r in rows if r['no']]
    if nos != list(range(nos[0], nos[-1] + 1)):
        bad.append('the numbers do not run on')
    for page in sorted({r['page'] for r in rows}):
        mine = [r for r in rows if r['page'] == page]
        body, latus = [r for r in mine if r['no']], [r for r in mine if not r['no']]
        if len(latus) != 1 or mine[-1] is not latus[0] or not latus[0]['name'].startswith('Latus'):
            bad.append('page %s: one Latus line, at the foot, is expected' % page)
            continue
        for r in body:
            if num(r['chr_m']) + num(r['chr_f']) != num(r['chr_sum']):
                bad.append('no. %s: the Christians do not add up' % r['no'])
            if num(r['jew_m']) + num(r['jew_f']) != num(r['jew_sum']):
                bad.append('no. %s: the Jews do not add up' % r['no'])
            if num(r['chr_sum']) + num(r['jew_sum']) != num(r['total']):
                bad.append('no. %s: the total is not the two sums' % r['no'])
        for k in FIGURES:
            s = sum(num(r[k]) for r in body)
            if s != num(latus[0][k]):
                bad.append('page %s, %s: the lines add up to %d, the Latus has %s' % (page, k, s, latus[0][k]))
    return bad


def cell(v):
    return '—' if v == '-' else v


def line(cells):
    out = io.StringIO()
    csv.writer(out, lineterminator='').writerow(cells)
    return out.getvalue()


def name_cell(r, kinds=None):
    kind = r['kind'] if kinds is None else kinds[r['kind']]
    return (kind + ' ' + r['name']).strip()


def read_source():
    """{(1, page id): [the lines of the page, as CSV]}: the heading opens the first page."""
    rows = table()
    parts = {}
    for r in rows:
        key = (1, PAGE_ID[r['page']])
        if not parts:
            parts[key] = [line(h) for h in HEAD]
        parts.setdefault(key, []).append(line([r['no'], name_cell(r)] + [cell(r[k]) for k in FIGURES]))
    return parts


def english():
    """The table in English, page by page: the heading and the kinds of place translated, the names as written."""
    def en_name(r):
        s = name_cell(r, KIND_EN)
        s = re.sub(r'(\S+)\[\?\]', r'[uncertain: \1]', s)        # the doubtful word, marked as elsewhere
        for old, new in (('d°', 'ditto'), ('dito', 'ditto')) + NAME_EN:
            s = s.replace(old, new)
        if s.startswith('Latus '):
            s = 'Page total ' + s[len('Latus '):]
        return s
    pages = {}
    for r in table():
        pid = PAGE_ID[r['page']]
        if not pages:
            pages[pid] = [line(h) for h in HEAD_EN]
        pages.setdefault(pid, []).append(line([r['no'], en_name(r)] + [cell(r[k]) for k in FIGURES]))
    return {1: ['\n'.join(pages[pid]) for pid in DOCS[0][1]]}


# ---------------------------------------------------------------- the spreadsheet

COLS = ['Page', 'No.', 'Kind of place', 'Name'] + [''] * 9 + ['Adds up?', 'Doubt (Claude)', 'Your remark',
                                                              'The line on the scan']
XLSX = 'table_check.xlsx'
FIRST = 8          # the first line of the table in the spreadsheet


def row_pictures(out_dir, scale=0.5, height=132):
    """A picture of each line, straightened: {(page, label): file}."""
    from PIL import Image
    os.makedirs(out_dir, exist_ok=True)
    where = json.load(io.open(os.path.join(HERE, 'rows.json'), encoding='utf-8'))
    made = {}
    for page, v in where.items():
        im = Image.open(os.path.join(RAW, v['file'])).convert('RGB')
        left, right, tilt = v['left'], v['right'], v['tilt']
        w = right - left
        for label, y in v['rows'].items():
            strip = Image.new('RGB', (w, height), 'white')
            for x in range(0, w, 100):
                yc = y + tilt * (x / w - 0.59)          # the line's height at this x
                top = int(round(yc - height / 2))
                strip.paste(im.crop((left + x, top, min(left + x + 100, right), top + height)), (x, 0))
            strip = strip.resize((int(w * scale), int(height * scale)), Image.LANCZOS)
            path = os.path.join(out_dir, '%s_%s.jpg' % (page, label))
            strip.save(path, quality=82)
            made[(page, label)] = path
    return made


def xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    rdir = os.path.join(ROOT, 'review', SLUG)
    pics = row_pictures(os.path.join(rdir, 'row_img'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Table'
    bold, thin = Font(bold=True), Side(style='thin', color='999999')
    ws['A1'] = '%s, "Die neuesten topografischen Nachrichten von Südpreussen": pages 401 to 406, entries 140 to 303' % REF
    ws['A1'].font = Font(bold=True, size=13)
    notes = [
        'Each line of the table is one line here, with a picture of that line of the scan at the right. Correct a cell '
        'where the scan has something else; write anything more under "Your remark". Save the file and tell Claude.',
        '"Adds up?" checks each line (men + women = sum, the two sums = total) and each "Latus" line against the lines '
        'of its page. It turns to LOOK when a figure you change breaks the sum. As read, every line and page adds up.',
        'A dash is the clerk\'s dash in an empty column. "[?]" marks a name not read with certainty; "Doubt" says why. '
        'The clerk\'s tailed stroke at the end of a shortened name is written as a period (or the colon he set).',
    ]
    for i, t in enumerate(notes, 2):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=False)
    # the printed heading, on three lines
    for i, h in enumerate(HEAD):
        for j, t in enumerate(h):
            if j < 2:               # "No." and the name column are set below: the name is two columns here
                continue
            c = ws.cell(FIRST - 3 + i, 3 + j, t)
            c.font, c.alignment = bold, Alignment(horizontal='center', vertical='center', wrap_text=True)
    top = FIRST - 3
    for a, b in (('E', 'E'), ('F', 'F'), ('M', 'M')):                # labels with nothing under them
        ws.merge_cells('%s%d:%s%d' % (a, top, b, top + 2))
    ws.merge_cells('G%d:L%d' % (top, top))
    ws.merge_cells('G%d:I%d' % (top + 1, top + 1))
    ws.merge_cells('J%d:L%d' % (top + 1, top + 1))
    # the name column of the page is two columns here
    ws.cell(top + 2, 3, 'Nahmen der Dörfer: kind of place (German script)').font = bold
    ws.cell(top + 2, 4, 'Nahmen der Dörfer: name').font = bold
    ws.cell(top + 2, 2, 'No.').font = bold
    ws.cell(top + 2, 1, 'Page').font = bold
    for j, t in ((14, 'Adds up?'), (15, 'Doubt (Claude)'), (16, 'Your remark'), (17, 'The line on the scan')):
        ws.cell(top + 2, j, t).font = bold
    for j in range(1, 18):
        ws.cell(top + 2, j).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(top + 2, j).border = Border(bottom=Side(style='medium'))
    ws.row_dimensions[top + 2].height = 45

    doubt, latus_fill = PatternFill('solid', fgColor='FFF2CC'), PatternFill('solid', fgColor='E7E6E6')
    r0 = {}                         # page -> first line
    n = FIRST
    for r in rows:
        r0.setdefault(r['page'], n)
        ws.cell(n, 1, int(r['page']))
        if r['no']:
            ws.cell(n, 2, int(r['no']))
        ws.cell(n, 3, r['kind'])
        ws.cell(n, 4, r['name'])
        for j, k in enumerate(FIGURES):
            ws.cell(n, 5 + j, '—' if r[k] == '-' else int(r[k])).alignment = Alignment(horizontal='right')
        if r['no']:
            f = ('=IF(AND(N(G{n})+N(H{n})=N(I{n}),N(J{n})+N(K{n})=N(L{n}),N(I{n})+N(L{n})=N(M{n})),"ok","LOOK")'
                 .format(n=n))
        else:
            a, b = r0[r['page']], n - 1
            f = '=IF(AND(%s),"ok","LOOK")' % ','.join('SUM({c}{a}:{c}{b})=N({c}{n})'.format(c=L(5 + j), a=a, b=b, n=n)
                                                     for j in range(9))
            for j in range(1, 16):
                ws.cell(n, j).font, ws.cell(n, j).fill = bold, latus_fill
                ws.cell(n, j).border = Border(top=thin)
        ws.cell(n, 14, f).alignment = Alignment(horizontal='center')
        if r['note']:
            c = ws.cell(n, 15, r['note'])
            c.fill, c.alignment = doubt, Alignment(wrap_text=True, vertical='center')
            ws.cell(n, 4).fill = doubt
        label = r['no'] or 'Latus'
        pic = XLImage(pics[(r['page'], label)])
        ws.add_image(pic, 'Q%d' % n)
        ws.row_dimensions[n].height = pic.height * 0.75 + 2
        for j in range(1, 17):
            ws.cell(n, j).alignment = Alignment(horizontal=ws.cell(n, j).alignment.horizontal,
                                                vertical='center', wrap_text=(j in (15, 16)))
        n += 1
    ws.conditional_formatting.add('N%d:N%d' % (FIRST, n - 1),
                                  CellIsRule(operator='equal', formula=['"LOOK"'],
                                             fill=PatternFill('solid', bgColor='FF7C80'), font=Font(bold=True)))
    widths = [6, 6, 12, 26] + [9] * 9 + [9, 44, 36, 180]
    for j, wd in enumerate(widths, 1):
        ws.column_dimensions[L(j)].width = wd
    ws.freeze_panes = 'E%d' % FIRST
    path = os.path.join(rdir, XLSX)
    wb.save(path)
    print('wrote', path, 'with', n - FIRST, 'lines')


def take(path, apply):
    """The editor's spreadsheet against table.csv: what they changed, and what they wrote."""
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    rows = table()
    got = []
    for n in range(FIRST, ws.max_row + 1):
        v = [ws.cell(n, j).value for j in range(1, 17)]
        if v[0] is None and v[3] is None:
            continue
        def s(x):
            return '' if x is None else str(x).strip()
        fig = ['-' if s(x) in ('—', '-', '') else str(int(float(s(x)))) for x in v[4:13]]
        got.append(dict(zip(['page', 'no', 'kind', 'name'] + FIGURES + ['remark'],
                            [s(v[0]), s(v[1]), s(v[2]), s(v[3])] + fig + [s(v[15])])))
    assert len(got) == len(rows), 'the spreadsheet has %d lines, table.csv %d' % (len(got), len(rows))
    changed = 0
    for old, new in zip(rows, got):
        label = old['no'] or old['name']
        for k in ['page', 'no', 'kind', 'name'] + FIGURES:
            if old[k] != new[k]:
                print('  %-10s %-9s %r -> %r' % (label, k, old[k], new[k]))
                changed += 1
                if apply:
                    old[k] = new[k]
        if new['remark']:
            print('  %-10s remark    %s' % (label, new['remark']))
    print(changed, 'cell(s) differ from table.csv')
    bad = check(rows if apply else [dict(o, **{k: n[k] for k in FIGURES}) for o, n in zip(rows, got)])
    for b in bad:
        print('  does not add up:', b)
    if apply and changed:
        with io.open(os.path.join(HERE, 'table.csv'), 'w', encoding='utf-8', newline='') as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator='\n')
            w.writeheader()
            w.writerows(rows)
        print('table.csv written. Clear a doubt in its last column where the remark settles it; then --write, '
              '--english, --cache and rebuild.')


# ---------------------------------------------------------------- the summaries

S = {
 1: ('Sechs Seiten einer Tabelle der Dörfer aus den Akten der preußischen Verwaltung Südpreußens. Sie enthalten die Einträge 140 bis 303 eines alphabetischen Verzeichnisses, von Kurow bis Wierzbna; die genannten Orte, darunter Ląd mit seinem Kloster, Lądek und Rzgow, liegen um Konin. Unter einem gedruckten Kopf gibt jeder Eintrag die Art des Ortes (Dorf, Hauland, Mühle, Vorwerk, Kloster, Krug), den Namen, die Zahl der Häuser und der Kirchen sowie die Christen und die Juden an, jeweils nach männlichem und weiblichem Geschlecht gezählt und summiert. Jede Seite schließt mit ihrer eigenen Summe, überschrieben „Latus“ und von 6 bis 11 gezählt. Trąbczyn, geschrieben „Trąpczyn“, ist Eintrag 291: 32 Häuser, 1 Kirche und 165 Einwohner, 51 männlich und 114 weiblich. Sein Hauland folgt als Eintrag 292 mit 34 Häusern und 271 Einwohnern. Szetlew hat 11 Häuser und 84 Einwohner, Szetlewek 20 Häuser und 140 Einwohner, darunter 4 Juden, und Łukom 28 Häuser und 149 Einwohner. Die Seiten tragen kein Datum und nennen den Kreis nicht.',
     'Six pages of a table of villages from the files of the Prussian administration of South Prussia. They hold entries 140 to 303 of an alphabetical list, from Kurow to Wierzbna; the places named, among them Ląd with its monastery, Lądek and Rzgow, lie around Konin. Under a printed heading each entry gives the kind of place (village, Hollander settlement, mill, demesne farm, monastery, inn), its name, the number of houses and of churches, and the Christians and the Jews, each counted as male and female and summed. Each page closes with its own total, headed "Latus" and numbered 6 to 11. Trąbczyn, written "Trąpczyn", is entry 291: 32 houses, 1 church and 165 inhabitants, 51 male and 114 female. Its Hollander settlement follows as entry 292 with 34 houses and 271 inhabitants. Szetlew has 11 houses and 84 inhabitants, Szetlewek 20 houses and 140 inhabitants, 4 of them Jews, and Łukom 28 houses and 149 inhabitants. The pages carry no date and do not name the district.'),
}
HOW =('written in the working session from the table as read (intake/holding.py, 2026-10-13).')


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    arg = sys.argv[1:]
    rows = table()
    bad = check(rows)
    print(len(rows), 'lines;', sum(1 for r in rows if r['no']), 'entries;', sum(1 for r in rows if '[?]' in r['name']),
          'marked [?];', sum(1 for r in rows if r['note']), 'with a note;', 'all sums hold' if not bad else 'PROBLEMS')
    for b in bad:
        print('  ', b)
    if '--crop' in arg:
        courtbook.cut(RAW, courtbook.placed(HERE, SCANS))
    if '--sheet' in arg:
        courtbook.fold_sheet(ROOT, SLUG, REF, RAW, courtbook.placed(HERE, SCANS), (), LEAF, note=FOLD_NOTE)
    if '--xlsx' in arg:
        xlsx(rows)
    if '--take' in arg:
        take(arg[arg.index('--take') + 1], '--apply' in arg)
    if '--write' in arg:
        assert not bad, 'the table does not add up'
        courtbook.write_pages(UNIT_DIR, DOCS, read_source())
    if '--english' in arg:
        tdir = os.path.join(HERE, 'translation')
        os.makedirs(tdir, exist_ok=True)
        courtbook.write_doc_yml(os.path.join(tdir, 'doc1.yml'), english()[1])
    if '--cache' in arg:
        courtbook.save_cache(ROOT, SLUG, os.path.join(HERE, 'translation'))
    if '--summaries' in arg:
        courtbook.write_summaries(ROOT, UNIT_DIR, SLUG, REF, S, HOW)


if __name__ == '__main__':
    main()
